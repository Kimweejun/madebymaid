"""PPT로 분석하기: deck.json -> .pptx + .html (same design), then publish to GitHub Pages and outputs/렘|람.

  python deck.py build   <deck.json>   # makes <name>.pptx and <name>.html next to the json
  python deck.py publish <deck.json>   # build + copy to outputs + git push + print the web address

Design: A4AI 발표 디자인 시스템 (1920×1080, Noto Sans KR, 흰 바탕, 위쪽 선 box·band, 로고 없음).
Colors: each deck picks one theme set (THEMES) that fits the report; the whole deck uses it.
Web page: the real PPT slides (rendered by PowerPoint) with each slide's 핵심 요약 beside it, outside the slide.
Rules: data/skills/madebymaid-ppt-design/SKILL.md · content: madebymaid-ppt-summary.
"""
import base64
import html
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import worklog  # noqa: E402  (komow-local/worklog.py)

ROOT = Path(__file__).resolve().parents[2]          # 통합서버
REPO = ROOT / 'madebymaid-reports'
PAGES = 'https://kimweejun.github.io/madebymaid-reports'
OUTBOX = ROOT / 'outputs'
MAIDS = {'렘', '람'}
MAID_BY = {'렘': 'rem', '람': 'ram'}

# A4AI 발표 디자인 시스템 tokens (tokens.json)
INK, PAPER, FIELD, LINE, MUTED, BODY2, NOTE, ON_DARK = '#17201A', '#FFFFFF', '#EEF2EB', '#D6DED2', '#5C685F', '#2C362F', '#7D897F', '#D5DDD2'
FONT = 'Noto Sans KR'
# Theme sets: one per deck, chosen by what the report is about. Each is a harmonized set of three colors
#   main  = lines, list dots, tags, band label (text-safe 4.5:1 on paper and on soft)
#   point = agenda numbers, cover accent      dark = cover background
THEMES = {
    '현장': {'main': '#3F5A3A', 'soft': '#DCE5D6', 'point': '#D9971E', 'dark': '#17201A'},  # 국방·제조·현장 (디자인 시스템 기본)
    '기술': {'main': '#2F4A6B', 'soft': '#DCE4EE', 'point': '#D9971E', 'dark': '#141C26'},  # 모델·도구·인프라
    '데이터': {'main': '#255B55', 'soft': '#D7EAE6', 'point': '#D9971E', 'dark': '#11201E'},  # 성능·지표·비교
    '사업': {'main': '#8E3427', 'soft': '#F3DDD8', 'point': '#D9971E', 'dark': '#221714'},  # 시장·비용·지원사업
    '연구': {'main': '#5E3F61', 'soft': '#EADFEA', 'point': '#4FA39A', 'dark': '#1D1720'},  # 논문·리서치·트렌드
}
VERDICTS = ('추천', '보류', '제외')
U = 1 / 144  # inches per design px (1920 px = 13.333 in)


def load(path):
    d = json.loads(Path(path).read_text(encoding='utf-8'))
    errors = []
    if not re.fullmatch(r'EXP-\d{4}-\d{2}', d.get('no', '')):
        errors.append('no: EXP-MMDD-NN 형식이어야 해요')
    if d.get('maid') not in MAIDS:
        errors.append('maid: 렘 또는 람')
    if not d.get('title'):
        errors.append('title이 비었어요')
    if d.get('theme') not in THEMES:
        errors.append(f'theme: {"·".join(THEMES)} 중 하나')
    slides = d.get('slides', [])
    if not 5 <= len(slides) <= 9:
        errors.append(f'본문 슬라이드는 5~9장이어야 해요 (표지·목차·결론 포함 8~12장). 지금 {len(slides)}장')
    for n, s in enumerate(slides, 3):
        if not s.get('title') or not s.get('points'):
            errors.append(f'{n}번 슬라이드: title과 points가 필요해요')
        if len(s.get('points', [])) > 6:
            errors.append(f'{n}번 슬라이드: points는 6개 이하')
        if not 1 <= len(s.get('summary', [])) <= 3:
            errors.append(f'{n}번 슬라이드: 핵심 요약(summary)은 1~3줄')
        errors += [f'{n}번 슬라이드 visual: {m}' for m in visual_errors(s.get('visual'), Path(path).parent)]
    need, have = len(slides) // 3, sum(1 for s in slides if s.get('visual'))
    if have < need:
        errors.append(f'본문 {len(slides)}장이면 그래프·그림이 {need}개 이상 필요해요 (지금 {have}개). '
                      '숫자는 그래프로, 기사·자료 사진은 image.file로 넣고, 맞는 사진이 없으면 image.prompt로 생성할 그림을 적어 주세요')
    c = d.get('conclusion', {})
    if c.get('verdict') not in VERDICTS:
        errors.append('conclusion.verdict: 추천·보류·제외 중 하나')
    if not c.get('headline') or not c.get('reasons'):
        errors.append('conclusion.headline과 reasons가 필요해요')
    if errors:
        worklog.log(d.get('no') or 'deck', MAID_BY.get(d.get('maid'), 'server'), 'error', 'deck.json: ' + '; '.join(errors)[:250])
        sys.exit('deck.json 확인 필요:\n- ' + '\n- '.join(errors))
    d['slug'] = re.sub(r'[^0-9A-Za-z가-힣]+', '-', d['title']).strip('-')[:40] or 'report'
    d['t'] = THEMES[d['theme']]
    return d


CHARTS = {'bar': '막대', 'compare': '비교 막대', 'pie': '파이', 'line': '꺾은선'}


def visual_errors(v, folder):
    """One visual per slide: a chart (bar·compare·pie·line) or an image (article photo, or a prompt to generate one)."""
    if not v:
        return []
    kind = v.get('type')
    if kind == 'image':
        if v.get('file'):
            f = folder / v['file']
            if not f.is_file():
                return [f'그림 파일이 없어요: {f}']
            v['path'] = f
        elif not v.get('prompt'):
            return ['image는 file(사진 파일) 또는 prompt(생성할 그림 설명)가 필요해요']
        errs = [] if v.get('caption') else ['caption(그림 해설, 예: "본 사진은 ~ 모습입니다")이 필요해요']
        if v.get('file') and not v.get('generated') and not v.get('source'):
            errs.append('기사·자료 사진은 source(출처)가 필요해요')
        return errs
    if kind not in CHARTS:
        return ['type은 bar·compare·pie·line·image 중 하나']
    labels, series = v.get('labels') or [], v.get('series') or []
    want = {'bar': (1, 1), 'pie': (1, 1), 'compare': (2, 2), 'line': (1, 3)}[kind]
    errs = []
    if not 2 <= len(labels) <= (6 if kind == 'pie' else 8):
        errs.append(f'labels는 2~{6 if kind == "pie" else 8}개')
    if not want[0] <= len(series) <= want[1]:
        errs.append(f'{CHARTS[kind]} 그래프의 series는 {want[0]}~{want[1]}개' if want[0] != want[1] else f'{CHARTS[kind]} 그래프의 series는 {want[0]}개')
    for sr in series:
        vals = sr.get('values') or []
        if len(vals) != len(labels) or not all(isinstance(x, (int, float)) for x in vals):
            errs.append(f'series "{sr.get("name", "")}"의 values는 labels 수만큼의 숫자')
    if not v.get('title'):
        errs.append('title(그래프 제목)이 필요해요')
    return errs


def reason_text(r):
    refs = r.get('slides') or []
    return r['text'] + (f' (근거: {", ".join(str(n) + "번" for n in refs)})' if refs else '')


def pages(d):
    """Every slide as a dict. `summary` is NOT drawn on the slide: the web page shows it beside the slide."""
    c = d['conclusion']
    out = [{'kind': 'cover', 'title': d['title'], 'summary': [d.get('subtitle') or d['title'], f'판정: {c["verdict"]}']},
           {'kind': 'agenda', 'title': '이 보고서의 흐름', 'lead': d.get('subtitle', ''), 'items': [s['title'] for s in d['slides']] + ['결론'],
            'summary': [f'결론 먼저: {c["headline"]}']}]
    out += [{'kind': 'content', **{k: s.get(k, '') for k in ('title', 'lead', 'points', 'band', 'source', 'gloss', 'summary', 'visual')}} for s in d['slides']]
    out.append({'kind': 'conclusion', 'title': c['headline'], 'lead': f'판정: {c["verdict"]}', 'points': [reason_text(r) for r in c['reasons']],
                'next': c.get('next') or [], 'band': c.get('band', ''), 'summary': [f'판정: {c["verdict"]}', *(c.get('next') or [])[:2]]})
    for n, p in enumerate(out, 1):
        p['n'], p['total'] = n, len(out)
    return out


def wrap(text, size, width, bold=False):
    """PowerPoint breaks Korean mid-word; break at spaces ourselves (vertical tab = line break in python-pptx).
    ponytail: width is estimated per character class, not measured; fine for Noto Sans KR at slide sizes."""
    em = lambda ch: 1.0 if ord(ch) > 0x2E80 else 0.33 if ch == ' ' else 0.62 if ch.isupper() or ch.isdigit() else 0.52
    limit, lines, cur = width / size / (1.03 if bold else 1), [], ''
    for word in text.split(' '):
        trial = (cur + ' ' + word) if cur else word
        if cur and sum(map(em, trial)) > limit:
            lines.append(cur); cur = word
        else:
            cur = trial
    return ''.join(lines + [cur])


# ── PowerPoint (1 design px = 1/144 in, font pt = px / 2; A4AI 발표 디자인 시스템 layout) ──
def build_pptx(d, out):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches, Pt

    t = d['t']
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(1920 * U), Inches(1080 * U)
    rgb = lambda h: RGBColor.from_string(h.lstrip('#'))
    px = lambda v: Inches(v * U)

    def rect(s, x, y, w, h, fill, round_=False):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
        if round_:
            sh.adjustments[0] = min(0.5, 12 / min(w, h))
        sh.shadow.inherit = False
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
        sh.line.fill.background()
        return sh

    def text(s, x, y, w, h, runs, size, color=INK, bold=False, gap=0, line=1.3, anchor=None):
        tb = s.shapes.add_textbox(px(x), px(y), px(w), px(h))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        if anchor:
            tf.vertical_anchor = anchor
        for i, line_ in enumerate(runs if isinstance(runs, list) else [runs]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after, p.line_spacing = Pt(gap / 2), line
            if isinstance(line_, str):
                line_ = [(wrap(line_, size, w, bold), color, bold)]
            for part, c, b, *sz in line_:
                for k, piece in enumerate(part.split('')):
                    if k:
                        p.add_line_break()
                    r = p.add_run(); r.text = piece
                    r.font.size, r.font.bold, r.font.name = Pt((sz[0] if sz else size) / 2), b, FONT
                    r.font.color.rgb = rgb(c)
        return tb

    def box(s, x, y, w, h, title, items):  # design-system box: 3px top line, title, dotted list
        rect(s, x, y, w, 3, t['main'])
        if title:
            text(s, x, y + 22, w, 50, title, 36, INK, True); y += 66
        tb = text(s, x, y + 22, w, h - 30, [[('■   ', t['main'], False, 16), (wrap(i, 28, w - 40), BODY2, False)] for i in items], 28, gap=14, line=1.4)
        for para in tb.text_frame.paragraphs:  # hanging indent: wrapped lines line up under the text, not the dot
            ppr = para._p.get_or_add_pPr(); ppr.set('marL', str(px(26))); ppr.set('indent', str(-px(26)))

    def chart(s, v, x, y, w, h):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        kind = v['type']
        cd = CategoryChartData(); cd.categories = v['labels']
        for sr in v['series']:
            cd.add_series(sr.get('name', ''), sr['values'])
        ctype = {'bar': XL_CHART_TYPE.COLUMN_CLUSTERED, 'compare': XL_CHART_TYPE.BAR_CLUSTERED,
                 'pie': XL_CHART_TYPE.PIE, 'line': XL_CHART_TYPE.LINE_MARKERS}[kind]
        ch = s.shapes.add_chart(ctype, px(x), px(y), px(w), px(h), cd).chart
        ch.has_title = False  # the slide draws its own chart title
        ch.font.size, ch.font.name = Pt(11), FONT
        ch.font.color.rgb = rgb(BODY2)
        colors = [t['main'], t['point'], '#93A097', '#5C685F', '#B8C4B4', '#D6DED2']
        plot = ch.plots[0]
        plot.has_data_labels = True
        dl = plot.data_labels; dl.font.size = Pt(11); dl.font.color.rgb = rgb(INK)
        if v.get('format'):
            dl.number_format, dl.number_format_is_linked = v['format'], False
        if kind == 'pie':
            dl.show_percentage, dl.show_value = True, False
            for i, point in enumerate(plot.series[0].points):
                point.format.fill.solid(); point.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
        else:
            if kind != 'line':
                plot.gap_width = 60
            for i, sr in enumerate(plot.series):
                c = rgb(colors[i])
                if kind == 'line':
                    sr.format.line.color.rgb = c; sr.format.line.width = Pt(2.5)
                    sr.marker.format.fill.solid(); sr.marker.format.fill.fore_color.rgb = c; sr.smooth = False
                else:
                    sr.format.fill.solid(); sr.format.fill.fore_color.rgb = c
            va = ch.value_axis
            va.has_major_gridlines = True; va.major_gridlines.format.line.color.rgb = rgb(LINE)
            va.format.line.fill.background(); va.tick_labels.font.color.rgb = rgb(MUTED)
            ch.category_axis.format.line.color.rgb = rgb(LINE)
            if kind == 'compare':
                ch.category_axis.reverse_order = True  # first label on top, like the slide reads
        ch.has_legend = len(v['series']) > 1 or kind == 'pie'
        if ch.has_legend:
            ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
            ch.legend.font.size = Pt(11)

    def picture(s, v, x, y, w, h):
        from PIL import Image
        src = v['path']
        if src.suffix.lower() not in ('.png', '.jpg', '.jpeg', '.gif', '.bmp'):  # webp etc.: PowerPoint-safe copy
            png = src.with_suffix('.slide.png'); Image.open(src).convert('RGB').save(png); src = png
        iw, ih = Image.open(src).size
        pic = s.shapes.add_picture(str(src), px(x), px(y), px(w), px(h))
        over = (iw / ih) / (w / h)  # crop to fill the frame without stretching (object-fit: cover)
        if over > 1:
            pic.crop_left = pic.crop_right = (1 - 1 / over) / 2
        else:
            pic.crop_top = pic.crop_bottom = (1 - over) / 2

    def visual(s, v, x, y, w, h):
        cap = (v.get('caption') or '') if v['type'] == 'image' else ''
        src = v.get('source') or ''
        foot = (34 if cap else 0) + (28 if src else 0)
        if v['type'] == 'image':
            if v.get('path'):
                picture(s, v, x, y, w, h - foot - 10)
            else:  # no photo yet: the prompt waits for the image-generation step
                rect(s, x, y, w, h - foot - 10, t['soft'], True)
                text(s, x + 40, y + 40, w - 80, 50, '이미지 생성 대기', 30, t['main'], True)
                text(s, x + 40, y + 100, w - 80, h - foot - 160, v['prompt'], 20, MUTED)
            label = '개념 이미지 (AI 생성) · ' if v.get('generated') or not v.get('path') else ''
            text(s, x, y + h - foot, w, 34, label + cap, 17, MUTED)
        else:
            text(s, x, y, w, 40, v['title'] + (f' ({v["unit"]})' if v.get('unit') else ''), 24, INK, True)
            chart(s, v, x, y + 46, w, h - foot - 56)
        if src:
            text(s, x, y + h - 28, w, 28, '출처: ' + src, 16, NOTE)

    for p in pages(d):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        if p['kind'] == 'cover':
            rect(s, 0, 0, 1920, 1080, t['dark'])
            rect(s, 110, 300, 120, 6, t['point'])
            text(s, 110, 330, 1500, 40, f'{d["no"]} · PPT로 분석하기', 24, t['point'], True)
            text(s, 110, 390, 1600, 520, [[(wrap(d['title'], 88, 1600, True), '#FFFFFF', True)], [(wrap(d.get('subtitle', ''), 30, 1600), ON_DARK, False, 30)]], 88, gap=40, line=1.1)
            text(s, 110, 960, 1600, 30, f'{d["maid"]} · {d.get("source", "")}', 18, '#9AA59C')
            continue
        text(s, 76, 60, 1768, 70, p['title'], 50, INK, True)
        if p.get('lead'):
            text(s, 76, 134, 1768, 36, p['lead'], 23, MUTED)
        text(s, 76, 1032, 1400, 24, f'{d["no"]} · {d["title"]}', 15, MUTED)
        text(s, 1644, 1032, 200, 24, f'{p["n"]} / {p["total"]}', 15, MUTED).text_frame.paragraphs[0].alignment = 3  # right
        band_h = 84 if p.get('band') else 0
        bottom = 1000 - (36 if p.get('source') else 0) - (40 if p.get('gloss') else 0) - (band_h + 20 if band_h else 0)
        top = 200
        if p['kind'] == 'agenda':
            for i, item in enumerate(p['items']):
                y = top + 24 + i * 70
                text(s, 76, y, 90, 60, f'{i + 3:02d}', 40, t['point'], True)
                text(s, 176, y + 6, 1600, 60, item, 30, INK, True)
        elif p['kind'] == 'conclusion':
            tag = rect(s, 76, top, 160, 56, t['main'], True)
            tf = tag.text_frame; tf.text = d['conclusion']['verdict']; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = tf.paragraphs[0]; para.alignment = 2  # center
            r = para.runs[0]; r.font.size, r.font.bold, r.font.name = Pt(14), True, FONT; r.font.color.rgb = rgb('#FFFFFF')
            box(s, 76, top + 86, 874, bottom - top - 86, '근거', p['points'])
            box(s, 970, top + 86, 874, bottom - top - 86, '다음 할 일', p['next'] or ['-'])
        elif p.get('visual'):  # points left, one chart or picture right
            box(s, 76, top, 700, bottom - top, '', p['points'])
            visual(s, p['visual'], 816, top, 1028, bottom - top)
        else:
            box(s, 76, top, 1768, bottom - top, '', p['points'])
        y = bottom + 20
        if band_h:
            rect(s, 76, y, 1768, band_h, t['soft'], True)
            text(s, 108, y, 1700, band_h, [[('핵심  ', t['main'], True), (p['band'], INK, False)]], 28, anchor=MSO_ANCHOR.MIDDLE)
            y += band_h + 12
        if p.get('source'):
            text(s, 76, y, 1768, 30, '출처: ' + p['source'], 16, NOTE); y += 36
        if p.get('gloss'):
            text(s, 76, y, 1768, 34, p['gloss'], 17, MUTED)
    prs.save(out)


def export_images(pptx):
    """PowerPoint itself renders each slide to PNG, so the web page shows the real slides. [] when PowerPoint is not there."""
    folder = pptx.parent / (pptx.stem + '_slides')
    shutil.rmtree(folder, ignore_errors=True); folder.mkdir()
    ps = (f"$pp = New-Object -ComObject PowerPoint.Application; $p = $pp.Presentations.Open('{pptx}', $true, $false, $false); "
          f"$i = 1; foreach ($s in $p.Slides) {{ $s.Export('{folder}\\' + $i + '.png', 'PNG', 1600, 900); $i++ }}; $p.Close(); $pp.Quit()")
    try:
        subprocess.run(['powershell', '-NoProfile', '-NonInteractive', '-Command', ps], capture_output=True, timeout=180)
    except (OSError, subprocess.TimeoutExpired):
        pass
    images = sorted(folder.glob('*.png'), key=lambda f: int(f.stem))
    return images


def slide_html(d, p):
    """Fallback when PowerPoint is not installed: the same slide drawn in HTML."""
    e = html.escape
    ul = lambda xs: '<ul>' + ''.join(f'<li>{e(x)}</li>' for x in xs) + '</ul>'
    if p['kind'] == 'cover':
        return (f'<section class="slide dark"><div class="cover"><span class="label">{e(d["no"])} · PPT로 분석하기</span><h1>{e(d["title"])}</h1>'
                f'<p>{e(d.get("subtitle", ""))}</p><small>{e(d["maid"])} · {e(d.get("source", ""))}</small></div></section>')
    if p['kind'] == 'agenda':
        body = '<ol class="toc">' + ''.join(f'<li><b>{n:02d}</b>{e(x)}</li>' for n, x in enumerate(p['items'], 3)) + '</ol>'
    elif p['kind'] == 'conclusion':
        body = (f'<div><span class="tag">{e(d["conclusion"]["verdict"])}</span><div class="two"><div class="box"><h3>근거</h3>{ul(p["points"])}</div>'
                f'<div class="box"><h3>다음 할 일</h3>{ul(p["next"] or ["-"])}</div></div></div>')
    else:
        v = p.get('visual')
        body = f'<div class="box">{ul(p["points"])}</div>' + (f'<div class="box"><h3>{e(v.get("title") or v.get("caption", ""))}</h3><p>그래프·그림은 PPT 파일에서 보세요.</p></div>' if v else '')
    return (f'<section class="slide"><div class="pad"><div class="hd"><h2>{e(p["title"])}' + (f'<span class="lead">{e(p["lead"])}</span>' if p.get('lead') else '')
            + f'</h2></div><div class="body">{body}</div>' + (f'<div class="band"><b>핵심</b>{e(p["band"])}</div>' if p.get('band') else '')
            + (f'<div class="src">출처: {e(p["source"])}</div>' if p.get('source') else '') + (f'<div class="gloss">{e(p["gloss"])}</div>' if p.get('gloss') else '')
            + f'</div><div class="foot"><span>{e(d["no"])} · {e(d["title"])}</span><span>{p["n"]} / {p["total"]}</span></div></section>')


# ── Web page: each PPT slide with its 핵심 요약 beside it (outside the slide) ──
def build_html(d, out, images):
    e, t = html.escape, d['t']
    rows = []
    for p in pages(d):
        img = images[p['n'] - 1] if len(images) == p['total'] else None
        slide = (f'<img src="data:image/png;base64,{base64.b64encode(img.read_bytes()).decode()}" alt="{p["n"]}번 슬라이드: {e(p["title"])}">'
                 if img else f'<div class="live">{slide_html(d, p)}</div>')
        rows.append(f'<article class="row" id="s{p["n"]}"><figure class="frame">{slide}</figure><aside><span class="no">{p["n"]} / {p["total"]}</span>'
                    f'<h3>핵심 요약</h3>' + ''.join(f'<p>{e(x)}</p>' for x in p['summary']) + '</aside></article>')
    pptx_name = f'{d["no"]}_{d["slug"]}.pptx'
    out.write_text(f'''<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{e(d["title"])} · {e(d["no"])}</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Noto+Sans+KR:wght@400;500;700;900&display=swap">
<style>
:root{{--ink:{INK};--paper:{PAPER};--line:{LINE};--muted:{MUTED};--body-2:{BODY2};--note:{NOTE};--on-dark:{ON_DARK};--main:{t["main"]};--soft:{t["soft"]};--point:{t["point"]};--dark:{t["dark"]}}}
*{{box-sizing:border-box}}body{{margin:0;background:#F3F4F2;color:var(--ink);font-family:"Noto Sans KR","Malgun Gothic",system-ui,sans-serif}}
.top{{position:sticky;top:0;z-index:5;display:flex;gap:16px;align-items:center;padding:10px 16px;background:var(--paper);border-bottom:3px solid var(--main);font-size:14px}}
.top span{{flex:1;min-width:0;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}}.top a{{color:var(--ink);font-weight:700}}
main{{max-width:1480px;margin:0 auto;padding:20px 16px}}
.row{{display:grid;grid-template-columns:minmax(0,1fr) 300px;gap:20px;align-items:start;margin:0 0 24px;scroll-margin-top:64px}}
.frame{{margin:0;aspect-ratio:16/9;position:relative;overflow:hidden;border-radius:6px;box-shadow:0 0 0 1px var(--line);background:var(--paper)}}
.frame img{{display:block;width:100%;height:100%}}
aside{{background:var(--paper);border-top:3px solid var(--main);padding:16px 18px 6px}}
aside .no{{font-size:13px;color:var(--muted)}}aside h3{{margin:6px 0 10px;font-size:17px;font-weight:700;color:var(--main)}}
aside p{{margin:0 0 10px;font-size:15px;line-height:1.55;color:var(--body-2)}}
@media (max-width:900px){{.row{{grid-template-columns:1fr;gap:8px}}}}
@media print{{.top{{display:none}}body{{background:#fff}}.row{{break-inside:avoid}}}}
/* fallback slides drawn in HTML (no PowerPoint) */
.live{{position:absolute;inset:0;container-type:inline-size}}
.slide{{position:absolute;left:0;top:0;width:1920px;height:1080px;background:var(--paper);transform-origin:0 0;overflow:hidden}}
.pad{{position:absolute;inset:0;padding:60px 76px 56px;display:flex;flex-direction:column}}
.hd{{margin-bottom:28px}}.hd h2{{margin:0;font-size:50px;font-weight:900;letter-spacing:-.02em;line-height:1.15}}
.hd .lead{{display:block;margin-top:10px;font-size:23px;font-weight:500;color:var(--muted)}}.body{{flex:1;min-height:0}}
.box{{border-top:3px solid var(--main);padding:22px 2px 6px}}.box h3{{margin:0 0 12px;font-size:36px}}
.box ul{{margin:0;padding:0;list-style:none;display:grid;gap:14px}}.box li{{padding-left:24px;position:relative;font-size:28px;line-height:1.5;color:var(--body-2)}}
.box li::before{{content:"";position:absolute;left:2px;top:.62em;width:8px;height:8px;border-radius:2px;background:var(--main)}}
.two{{display:grid;grid-template-columns:1fr 1fr;gap:20px;margin-top:30px}}
.tag{{display:inline-block;font-size:28px;font-weight:700;padding:8px 26px;border-radius:999px;background:var(--main);color:#fff}}
.band{{margin-top:20px;display:flex;gap:18px;border-radius:12px;padding:22px 32px;background:var(--soft);font-size:28px}}.band b{{color:var(--main)}}
.src{{font-size:16px;color:var(--note);margin-top:12px}}.gloss{{margin-top:10px;font-size:17px;color:var(--muted)}}
.foot{{position:absolute;left:76px;right:76px;bottom:22px;display:flex;justify-content:space-between;font-size:15px;color:var(--muted)}}
.toc{{margin:0;padding:24px 0 0;list-style:none;display:grid;gap:22px;font-size:30px;font-weight:700}}.toc b{{display:inline-block;width:100px;font-size:40px;font-weight:900;color:var(--point)}}
.dark{{background:var(--dark);color:#fff}}.cover{{position:absolute;left:110px;right:110px;top:50%;transform:translateY(-50%)}}
.cover .label{{display:block;padding-top:24px;border-top:6px solid var(--point);width:max-content;font-size:24px;font-weight:700;color:var(--point)}}
.cover h1{{margin:30px 0 24px;font-size:88px;font-weight:900;line-height:1.1}}.cover p{{margin:0;font-size:30px;color:var(--on-dark)}}.cover small{{display:block;margin-top:60px;font-size:18px;color:#9AA59C}}
</style></head><body><div class="top"><span><b>{e(d["no"])}</b> {e(d["title"])}</span><a href="{e(pptx_name)}" download>PPT 받기</a><a href="../">보고서 목록</a></div>
<main>{"".join(rows)}</main>
<script>
const fit=()=>document.querySelectorAll('.live').forEach(l=>l.firstElementChild.style.transform='scale('+l.clientWidth/1920+')');addEventListener('resize',fit);fit();
const rows=[...document.querySelectorAll('.row')];
addEventListener('keydown',ev=>{{const k=ev.key;if(!['ArrowRight','ArrowLeft','PageDown','PageUp'].includes(k))return;ev.preventDefault();
const mid=scrollY+innerHeight/2;let i=rows.findIndex(r=>r.offsetTop+r.offsetHeight>mid);if(i<0)i=rows.length-1;
i=Math.max(0,Math.min(rows.length-1,i+(k==='ArrowLeft'||k==='PageUp'?-1:1)));rows[i].scrollIntoView({{behavior:'smooth',block:'start'}});}});
</script></body></html>''', encoding='utf-8')


def build(path):
    d = load(path)
    who = MAID_BY.get(d['maid'], 'server')
    worklog.log(d['no'], who, 'step', 'deck: build')
    base = Path(path).parent / f'{d["no"]}_{d["slug"]}'
    pptx = base.with_suffix('.pptx')
    build_pptx(d, pptx)
    build_html(d, base.with_suffix('.html'), export_images(pptx.resolve()))
    prompts = [(n, sl['visual']['prompt']) for n, sl in enumerate(d['slides'], 3)
               if (sl.get('visual') or {}).get('type') == 'image' and not sl['visual'].get('file')]
    pending = Path(path).parent / f'{d["no"]}_이미지프롬프트.md'
    # always written: the experiment card shows or clears "그림 생성 필요" from it (image-gen skill fills the prompts)
    pending.write_text(f'# {d["no"]} 생성할 그림 (테마: {d["theme"]})\n\n' + (''.join(f'## {n}번 슬라이드\n\n{pr}\n\n' for n, pr in prompts) if prompts else '생성할 그림 없음\n'),
                       encoding='utf-8')
    return d, base


def git(*args):
    r = subprocess.run(['git', '-C', str(REPO), *args], capture_output=True, text=True, encoding='utf-8')
    if r.returncode:
        worklog.log('deck', 'server', 'error', f'git {args[0]} failed: {r.stderr.strip()[:200]}')
        sys.exit(f'git {" ".join(args)} 실패: {r.stderr.strip()}')
    return r.stdout


def write_index():
    rows = []
    for meta in sorted(REPO.glob('EXP-*/meta.json'), reverse=True):
        m = json.loads(meta.read_text(encoding='utf-8'))
        rows.append(f'<li><a href="{m["no"]}/"><b>{html.escape(m["no"])}</b> {html.escape(m["title"])}</a> <span>{html.escape(m["verdict"])} · {html.escape(m["maid"])}</span></li>')
    (REPO / 'index.html').write_text(
        '<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>메이드 바이 메이드 보고서</title>'
        '<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Noto+Sans+KR:wght@400;700;900&display=swap">'
        f'<style>body{{font-family:"Noto Sans KR","Malgun Gothic",sans-serif;max-width:860px;margin:40px auto;padding:0 16px;color:{INK}}}'
        f'h1{{font-weight:900;padding-bottom:12px;border-bottom:3px solid {INK}}}li{{margin:12px 0}}a{{color:{INK};text-decoration:none}}a:hover{{text-decoration:underline}}'
        f'span{{color:{MUTED};font-size:14px}}</style></head>'
        f'<body><h1>메이드 바이 메이드 실험 보고서</h1><ul>{"".join(rows) or "<li>아직 올린 보고서가 없어요.</li>"}</ul></body></html>', encoding='utf-8')


def publish(path):
    d, base = build(path)
    if not (REPO / '.git').is_dir():
        sys.exit(f'보고서 저장소가 없어요: {REPO}')
    git('pull', '-q', '--rebase')
    folder = REPO / d['no']
    folder.mkdir(exist_ok=True)
    shutil.copy(base.with_suffix('.html'), folder / 'index.html')
    shutil.copy(base.with_suffix('.pptx'), folder / base.with_suffix('.pptx').name)
    (folder / 'meta.json').write_text(json.dumps({k: d[k] for k in ('no', 'title', 'maid')} | {'verdict': d['conclusion']['verdict']}, ensure_ascii=False), encoding='utf-8')
    write_index()
    git('add', '-A')
    if git('status', '--porcelain').strip():
        git('commit', '-q', '-m', f'{d["no"]} {d["title"]}')
    git('push', '-q', 'origin', 'HEAD')
    url = f'{PAGES}/{d["no"]}/'
    worklog.log(d['no'], MAID_BY.get(d['maid'], 'server'), 'done', 'deck published', url)
    out = OUTBOX / d['maid']
    out.mkdir(parents=True, exist_ok=True)
    for ext in ('.pptx', '.html'):
        shutil.copy(base.with_suffix(ext), out / base.with_suffix(ext).name)
    (out / f'{d["no"]}_웹주소.txt').write_text(url + '\n', encoding='utf-8')
    pending = Path(path).parent / f'{d["no"]}_이미지프롬프트.md'
    if pending.exists():
        shutil.copy(pending, out / pending.name)
    print(url)
    return url


if __name__ == '__main__':
    if len(sys.argv) != 3 or sys.argv[1] not in ('build', 'publish'):
        sys.exit(__doc__)
    if sys.argv[1] == 'build':
        print(build(sys.argv[2])[1])
    else:
        publish(sys.argv[2])
